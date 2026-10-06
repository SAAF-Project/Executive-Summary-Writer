"""Tests for reading the audit material by input source type. No API calls are made.

pytest tests/
"""

import shutil
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import material  # noqa: E402
from material import (  # noqa: E402
    DEFAULT_SOURCE_TYPE,
    default_source,
    load_material,
)


def text_of(blocks):
    assert len(blocks) == 1 and blocks[0]["type"] == "text"
    return blocks[0]["text"]


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Findings"
    table = slide.shapes.add_table(
        1, 2, Inches(1), Inches(2), Inches(6), Inches(1)
    ).table
    table.cell(0, 0).text = "F1"
    table.cell(0, 1).text = "Checks not performed"
    slide.notes_slide.notes_text_frame.text = "Agreed with management"
    presentation.save(str(path))


def make_xlsx(path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Findings"
    sheet.append(["F1", "Checks not performed", 3])
    workbook.save(str(path))


def make_docx(path):
    from docx import Document

    document = Document()
    document.add_paragraph("Overall conclusion: needs improvement")
    document.save(str(path))


def test_text_types_are_read_as_they_are(tmp_path):
    for suffix in ("txt", "md"):
        path = tmp_path / f"report.{suffix}"
        path.write_text("Finding F1", encoding="utf-8")
        assert "Finding F1" in text_of(load_material([path]))


def test_pptx_gives_slide_text_tables_and_notes(tmp_path):
    path = tmp_path / "report.pptx"
    make_pptx(path)

    text = text_of(load_material([path]))
    assert "--- Slide 1 ---" in text
    assert "Findings" in text
    assert "F1 | Checks not performed" in text
    assert "Notes: Agreed with management" in text


def test_xlsx_gives_one_line_per_row(tmp_path):
    path = tmp_path / "report.xlsx"
    make_xlsx(path)

    text = text_of(load_material([path]))
    assert "--- Sheet Findings ---" in text
    assert "F1 | Checks not performed | 3" in text


def test_docx_is_read(tmp_path):
    path = tmp_path / "report.docx"
    make_docx(path)
    assert "needs improvement" in text_of(load_material([path]))


def test_given_type_wins_over_the_file_name(tmp_path):
    path = tmp_path / "upload.bin"
    make_pptx(path)
    assert "Checks not performed" in text_of(load_material([path], "pptx"))


def test_wrong_type_is_reported(tmp_path):
    path = tmp_path / "report.md"
    path.write_text("Finding F1", encoding="utf-8")
    with pytest.raises(ValueError, match="Could not read report.md as .pptx"):
        load_material([path], "pptx")


def test_unsupported_extension_is_refused(tmp_path):
    path = tmp_path / "report.zip"
    path.write_bytes(b"x")
    with pytest.raises(ValueError, match="Unsupported file type"):
        load_material([path])


def test_legacy_type_without_libreoffice_says_what_to_do(tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: None)
    path = tmp_path / "report.ppt"
    path.write_bytes(b"x")
    with pytest.raises(ValueError, match="needs LibreOffice"):
        load_material([path])


def test_default_is_the_pptx_sample():
    assert DEFAULT_SOURCE_TYPE == "pptx"
    path = default_source(DEFAULT_SOURCE_TYPE)
    assert path.suffix == ".pptx"
    if path.is_file():  # the sample may not be part of every checkout
        assert "--- Slide 1 ---" in text_of(load_material([path], "pptx"))


def test_type_without_a_sample_needs_a_file():
    assert default_source("md") == material.DEFAULT_SOURCES["md"]
    with pytest.raises(ValueError, match="no sample of type xlsx"):
        default_source("xlsx")


def make_finding_deck(path):
    """A slide titled "Findings and recommendations" with the table of one finding, as in the template."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Findings and recommendations"
    table = slide.shapes.add_table(
        4, 3, Inches(1), Inches(2), Inches(8), Inches(3)
    ).table
    rows = [
        ("0123 - Access rights are not reviewed", "", "Risk: High"),
        ("Finding", "Access rights were last reviewed in 2023.", "AO02"),
        ("Root Cause(s)", "Nobody owns the review.", ""),
        ("Recommendation", "Appoint an owner and review each quarter.", ""),
    ]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    table.cell(0, 0).merge(table.cell(0, 1))  # the title spans the label and the content column
    table.cell(0, 1).text = "hidden under the title"
    presentation.save(str(path))


def test_finding_slide_is_read_as_labelled_plain_text(tmp_path):
    path = tmp_path / "report.pptx"
    make_finding_deck(path)

    text = text_of(load_material([path]))
    assert (
        "Finding title: 0123 - Access rights are not reviewed\n"
        "Risk: High\n"
        "Finding: Access rights were last reviewed in 2023. (AO02)\n"
        "Root Cause(s): Nobody owns the review.\n"
        "Recommendation: Appoint an owner and review each quarter."
    ) in text
    assert "hidden under the title" not in text
    assert " | " not in text
