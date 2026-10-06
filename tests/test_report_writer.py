"""Tests for writing the summary onto the report deck.

pytest tests/
"""

import logging
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from material import DEFAULT_SOURCES  # noqa: E402
from report_sections import (  # noqa: E402
    count_risks,
    finding_owner,
    read_finding_table,
    read_findings,
)
from report_writer import (  # noqa: E402
    KEY_FIGURES_PLACEHOLDER,
    fill_finding_numbers,
    fill_header,
    fill_grade,
    fill_findings_table,
    fill_positive_block,
    fill_summary_block,
    find_block,
    find_findings_table,
    find_summary_block,
    slide_text,
    write_summary_to_deck,
)

TEMPLATE = DEFAULT_SOURCES["pptx"]

SUMMARY = """\
### Executive Board Summary

Controls exist and are documented.
The overall conclusion is needs improvement.

Ownership is **divided**.

Management has agreed actions.

### Potential Gaps for Executive Board Consideration

- Financial effect not quantified.
"""

POS_POINTS = ["Payment is well controlled.", "Model is up-to-date."]
NEG_POINTS = {
    "Payment is not well controlled.": "The team should have a gateway.",
    "Model is not up-to-date.": "There should be a product owner.",
}


def make_deck(block_heading="Audit conclusion", finding_rows=4):
    """Cover, table of contents, and an executive summary slide with the blocks and the table to fill."""
    deck = Presentation()
    slides = [deck.slides.add_slide(deck.slide_layouts[6]) for _ in range(3)]
    icon = slides[1].shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(1), Inches(1)
    )
    icon.click_action.target_slide = slides[2]
    label = slides[1].shapes.add_textbox(
        Inches(1), Inches(2.2), Inches(3), Inches(0.5)
    )
    label.text_frame.text = "Executive summary"

    cell = (
        slides[2]
        .shapes.add_table(1, 1, Inches(1), Inches(1), Inches(6), Inches(2.5))
        .table.cell(0, 0)
    )
    cell.text_frame.text = block_heading
    placeholder = cell.text_frame.add_paragraph()
    for text in ("[AUDIT SUMMARY] + ", "[CONCLUSION]"):
        run = placeholder.add_run()
        run.text = text
        run.font.size = Pt(10)
        run.font.italic = True
    cell.text_frame.add_paragraph()

    cell = (
        slides[2]
        .shapes.add_table(1, 1, Inches(7), Inches(1), Inches(3), Inches(2.5))
        .table.cell(0, 0)
    )
    cell.text_frame.text = "Positive aspects"
    run = cell.text_frame.add_paragraph().add_run()
    run.text = "[POSITIVE OBSERVATIONS]"
    run.font.size = Pt(10)
    run.font.italic = True

    table = (
        slides[2]
        .shapes.add_table(
            finding_rows + 1, 3, Inches(1), Inches(4), Inches(9), Inches(2)
        )
        .table
    )
    for column, text in enumerate(
        ("Main findings", "Recommendation", "Finding owner")
    ):
        table.cell(0, column).text_frame.text = text
    for column, text in enumerate(
        ("[FINDINGS", "[RECOMMENDATION]", "[FINDING OWNER]")
    ):
        paragraph = table.cell(1, column).text_frame.paragraphs[0]
        for part in (text, " more]") if column == 0 else (text,):
            run = paragraph.add_run()
            run.text = part
            run.font.italic = True
    return deck


def findings_rows(deck):
    """The rows below the header of the findings table, as lists of cell texts."""
    _, table = find_findings_table(deck)
    return [[cell.text for cell in row.cells] for row in list(table.rows)[1:]]


def test_only_the_summary_paragraphs_go_on_the_slide():
    assert slide_text(SUMMARY) == [
        "Controls exist and are documented. The overall conclusion is needs improvement.",
        "Ownership is divided.",
        "Management has agreed actions.",
    ]


def test_summary_without_the_heading_is_refused():
    with pytest.raises(ValueError, match="no text under"):
        slide_text("### Something else\n\nText")


def test_placeholder_is_replaced_and_heading_and_format_are_kept():
    deck = make_deck()
    assert fill_summary_block(deck, SUMMARY) == 3

    _, _, cell = find_summary_block(deck)
    paragraphs = cell.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Audit conclusion"] + slide_text(
        SUMMARY
    )
    for paragraph in paragraphs[1:]:
        assert len(paragraph.runs) == 1
        assert paragraph.runs[0].font.italic is True
        assert paragraph.runs[0].font.size == Pt(10)


def test_long_summary_gets_a_smaller_font_and_then_a_warning(caplog):
    deck = make_deck()
    fill_summary_block(
        deck, f"### Executive Board Summary\n\n{'word ' * 300}"
    )
    size = find_summary_block(deck)[2].text_frame.paragraphs[1].runs[0].font.size
    assert Pt(7) <= size < Pt(10)

    with caplog.at_level(logging.WARNING):
        fill_summary_block(
            make_deck(), f"### Executive Board Summary\n\n{'word ' * 2000}"
        )
    assert "too long" in caplog.text


def test_positive_points_become_bullets_under_the_heading():
    deck = make_deck()
    assert fill_positive_block(deck, POS_POINTS) == 3

    _, _, cell = find_block(deck, "Positive aspects")
    paragraphs = cell.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Positive aspects"] + POS_POINTS
    # every point is a copy of the placeholder, which carries the bullet of the template
    assert all(p.runs[0].font.italic is True for p in paragraphs[1:])


def test_findings_get_a_row_each_and_recommendations_cite_the_model():
    deck = make_deck()
    assert fill_findings_table(deck, NEG_POINTS, "gpt-5.1") == 3

    assert findings_rows(deck) == [
        [
            "Payment is not well controlled.",
            "The team should have a gateway. (suggested by AI|gpt-5.1)",
            "[FINDING OWNER]",
        ],
        [
            "Model is not up-to-date.",
            "There should be a product owner. (suggested by AI|gpt-5.1)",
            "",
        ],
        ["", "", ""],
        ["", "", ""],
    ]
    _, table = find_findings_table(deck)
    for row in (1, 2):
        for column in (0, 1):
            runs = table.cell(row, column).text_frame.paragraphs[0].runs
            assert len(runs) == 1 and runs[0].font.italic is True


def test_long_findings_get_a_smaller_font_and_the_table_keeps_its_height(caplog):
    def font_and_height(findings):
        deck = make_deck()
        _, table = find_findings_table(deck)
        before = sum(row.height for row in table.rows)
        fill_findings_table(deck, findings, "gpt-5.1")
        sizes = {
            table.cell(row, column).text_frame.paragraphs[0].runs[0].font.size
            for row in range(1, len(findings) + 1)
            for column in (0, 1)
        }
        assert len(sizes) == 1  # one size for the whole table
        return sizes.pop(), [row.height for row in list(table.rows)[1:]], before

    def same_height(heights, before):
        """The height of the table is only shared out differently over its rows."""
        return abs(sum(heights) + Pt(28.8) - before) < Pt(1)  # plus the header row

    size, heights, before = font_and_height(NEG_POINTS)
    assert size == Pt(10) and same_height(heights, before)

    long = {f"Finding {n}. " + "word " * 16: "Do this. " + "word " * 8 for n in range(4)}
    size, heights, before = font_and_height(long)
    assert Pt(7) <= size < Pt(10) and same_height(heights, before)

    # a long first row takes its height from the rows that need less
    uneven = {"Finding. " + "word " * 40: "Fix it.", "Short finding.": "Fix it."}
    size, heights, before = font_and_height(uneven)
    assert heights[0] > 2 * heights[1] and same_height(heights, before)
    assert min(heights[2:]) >= Pt(19)  # the empty rows keep room for one line

    with caplog.at_level(logging.WARNING):
        size, _, _ = font_and_height(
            {f"Finding {n}. " + "word " * 200: "Fix it." for n in range(4)}
        )
    assert size == Pt(7) and "too long" in caplog.text


def test_findings_beyond_the_rows_of_the_table_are_reported(caplog):
    deck = make_deck(finding_rows=1)
    with caplog.at_level(logging.WARNING):
        fill_findings_table(deck, NEG_POINTS, "gpt-5.1")

    assert [row[0] for row in findings_rows(deck)] == [
        "Payment is not well controlled."
    ]
    assert "Not written: Model is not up-to-date." in caplog.text


def test_deck_without_the_block_is_refused():
    with pytest.raises(ValueError, match="No block starting with"):
        fill_summary_block(make_deck("Something else"), SUMMARY)


def test_a_copy_is_saved_and_the_input_is_left_alone(tmp_path):
    source = tmp_path / "report.pptx"
    make_deck().save(str(source))
    before = source.read_bytes()

    out_file = tmp_path / "out" / "report_summary.pptx"
    assert (
        write_summary_to_deck(
            source, SUMMARY, out_file, POS_POINTS, NEG_POINTS, "gpt-5.1"
        )
        == 3
    )
    assert source.read_bytes() == before
    saved = Presentation(str(out_file))
    _, _, cell = find_summary_block(saved)
    assert "Ownership is divided." in cell.text
    assert "Model is up-to-date." in find_block(saved, "Positive aspects")[2].text
    assert findings_rows(saved)[0][0] == "Payment is not well controlled."


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not in checkout")
def test_template(tmp_path):
    out_file = tmp_path / "template.pptx"
    assert write_summary_to_deck(TEMPLATE, SUMMARY, out_file) == 4

    deck = Presentation(str(out_file))
    _, _, cell = find_summary_block(deck)
    assert cell.text.startswith("Audit conclusion\nControls exist")
    assert "[AUDIT SUMMARY" not in cell.text
    # the rest of the slide is untouched
    slide = "\n".join(
        shape.table.cell(0, 0).text
        for shape in list(deck.slides)[3].shapes
        if getattr(shape, "has_table", False)
    )
    assert "[POSITIVE OBSERVATIONS]" in slide


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not in checkout")
def test_template_with_points(tmp_path):
    out_file = tmp_path / "template.pptx"
    write_summary_to_deck(
        TEMPLATE, SUMMARY, out_file, POS_POINTS, NEG_POINTS, "gpt-5.1"
    )

    deck = Presentation(str(out_file))
    _, _, cell = find_block(deck, "Positive aspects")
    assert cell.text == "\n".join(["Positive aspects"] + POS_POINTS)
    # the points keep the bullet of the placeholder
    for paragraph in cell.text_frame.paragraphs[1:]:
        assert paragraph._p.pPr.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"
        ) is not None
    rows = findings_rows(deck)
    assert rows[0] == [
        "Payment is not well controlled.",
        "The team should have a gateway. (suggested by AI|gpt-5.1)",
        "[FINDING OWNER]\n",
    ]
    assert rows[1][:2] == [
        "Model is not up-to-date.",
        "There should be a product owner. (suggested by AI|gpt-5.1)",
    ]
    assert rows[2][:2] == ["", ""]


# -------------------- Number of findings per risk level --------------------

RISKS = [
    ("0101 - Access rights are not reviewed", "Risk: High"),
    ("0102 - Reports contain errors", "Risk: Medium"),
    ("0103 - Actions are not tracked", "Risk: medium"),
    ("0104 - Not graded yet", "Risk: xxx"),
]


def make_numbers_deck(risks=RISKS, overview_rows=3):
    """A deck with the "Findings" box, the "Finding overview" table and one slide per finding."""
    deck = Presentation()
    blank, titled = deck.slide_layouts[6], deck.slide_layouts[5]
    cover, contents, summary, overview = (
        deck.slides.add_slide(blank) for _ in range(4)
    )
    for position, (name, target) in enumerate(
        (("Executive summary", summary), ("Finding overview", overview))
    ):
        left = Inches(1 + 4 * position)
        icon = contents.shapes.add_shape(
            MSO_SHAPE.OVAL, left, Inches(1), Inches(1), Inches(1)
        )
        icon.click_action.target_slide = target
        contents.shapes.add_textbox(
            left, Inches(2.2), Inches(3), Inches(0.5)
        ).text_frame.text = name

    box = summary.shapes.add_group_shape()
    for column, (letter, number) in enumerate(zip("HML", "126")):
        for row, text in enumerate((letter, number)):
            shape = box.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(6 + 0.5 * column),
                Inches(1.5 + 0.4 * row),
                Inches(0.4),
                Inches(0.3),
            )
            shape.text_frame.paragraphs[0].add_run().text = text
    heading = box.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6), Inches(1), Inches(1.4), Inches(0.3)
    )
    heading.text_frame.paragraphs[0].add_run().text = "Findings"
    square = box.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5), Inches(1), Inches(0.9), Inches(0.9)
    )
    square.text_frame.paragraphs[0].add_run().text = "B"

    table = overview.shapes.add_table(
        overview_rows + 2, 4, Inches(1), Inches(1), Inches(8), Inches(2)
    ).table
    for column, text in enumerate(("AO01. [AO DESCRIPTION]", "[X]", "[X]", "[X]")):
        table.cell(1, column).text_frame.paragraphs[0].add_run().text = text
    for column, text in enumerate(("Total", "[X]", "[X]", "[X]")):
        table.cell(overview_rows + 1, column).text_frame.paragraphs[0].add_run().text = text

    for title, risk in risks:
        slide = deck.slides.add_slide(titled)
        slide.shapes.title.text = "Findings and recommendations"
        finding = slide.shapes.add_table(
            2, 3, Inches(1), Inches(2), Inches(8), Inches(1)
        ).table
        for column, text in enumerate((title, "", risk)):
            finding.cell(0, column).text = text
        finding.cell(1, 0).text = "Finding"
        # the response to a finding repeats its title and must not be counted
        response = deck.slides.add_slide(titled)
        response.shapes.title.text = "Management Response"
        response.shapes.add_table(
            1, 2, Inches(1), Inches(2), Inches(8), Inches(1)
        ).table.cell(0, 0).text = title
    return deck


def overview_rows(deck):
    table = next(
        shape.table
        for shape in list(deck.slides)[3].shapes
        if getattr(shape, "has_table", False)
    )
    return [[cell.text for cell in row.cells] for row in list(table.rows)[1:]]


def test_title_and_risk_are_read_from_every_finding_slide():
    findings = read_findings(make_numbers_deck())

    assert findings == [
        ("0101 - Access rights are not reviewed", "High"),
        ("0102 - Reports contain errors", "Medium"),
        ("0103 - Actions are not tracked", "Medium"),
        ("0104 - Not graded yet", None),
    ]
    assert count_risks(findings) == [1, 2, 0]


def test_numbers_go_in_the_findings_box_and_the_overview(caplog):
    deck = make_numbers_deck()
    with caplog.at_level(logging.WARNING):
        fill_finding_numbers(deck)

    box = next(s for s in list(deck.slides)[2].shapes if hasattr(s, "shapes"))
    assert [s.text_frame.text for s in box.shapes] == [
        "H", "1", "M", "2", "L", "0", "Findings", "B",
    ]  # fmt: skip
    # one more finding than the table had rows: a row was added above the total
    assert overview_rows(deck) == [
        ["0101 - Access rights are not reviewed", "1", "", ""],
        ["0102 - Reports contain errors", "", "1", ""],
        ["0103 - Actions are not tracked", "", "1", ""],
        ["0104 - Not graded yet", "", "", ""],
        ["Total", "1", "2", "0"],
    ]
    assert "No risk level" in caplog.text and "0104 - Not graded yet" in caplog.text


def test_overview_rows_that_are_not_needed_are_emptied():
    deck = make_numbers_deck(risks=RISKS[:1], overview_rows=3)
    fill_finding_numbers(deck)

    assert overview_rows(deck) == [
        ["0101 - Access rights are not reviewed", "1", "", ""],
        ["", "", "", ""],
        ["", "", "", ""],
        ["Total", "1", "0", "0"],
    ]


def test_overview_of_a_report_that_is_filled_in_already_is_replaced():
    deck = make_numbers_deck(risks=RISKS[:2], overview_rows=3)
    table = next(
        s.table for s in list(deck.slides)[3].shapes if getattr(s, "has_table", False)
    )
    # as in a finished report: a mark only under one level per row, the other cells empty
    for column in (2, 3):
        table.cell(1, column).text_frame.paragraphs[0].runs[0]._r.getparent().remove(
            table.cell(1, column).text_frame.paragraphs[0].runs[0]._r
        )
    table.cell(1, 0).text_frame.paragraphs[0].runs[0].text = "AO01. Purchase approval"
    fill_finding_numbers(deck)

    assert overview_rows(deck) == [
        ["0101 - Access rights are not reviewed", "1", "", ""],
        ["0102 - Reports contain errors", "", "1", ""],
        ["", "", "", ""],
        ["Total", "1", "1", "0"],
    ]


def test_deck_without_finding_slides_is_left_as_it_is():
    deck = make_deck()
    assert fill_finding_numbers(deck) == []


def test_grade_goes_in_the_square_in_its_colour():
    deck = make_numbers_deck()
    assert fill_grade(deck, "C") == 3

    box = next(s for s in list(deck.slides)[2].shapes if hasattr(s, "shapes"))
    square = list(box.shapes)[-1]
    assert square.text_frame.text == "C"
    assert str(square.fill.fore_color.rgb) == "FFC000"
    with pytest.raises(ValueError, match="Unknown grade"):
        fill_grade(deck, "E")


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not in checkout")
def test_template_with_grade(tmp_path):
    out_file = tmp_path / "template.pptx"
    write_summary_to_deck(TEMPLATE, SUMMARY, out_file, grade="D")

    deck = Presentation(str(out_file))
    box = next(
        s for s in list(deck.slides)[3].shapes if hasattr(s, "shapes")
    )
    texts = [s.text_frame.text for s in box.shapes]
    assert "D" in texts and "B" not in texts
    square = next(s for s in box.shapes if s.text_frame.text == "D")
    assert str(square.fill.fore_color.rgb) == "FF0000"


# -------------------- Header of the executive summary slide --------------------


def make_header_deck():
    """Title slide, table of contents, and an executive summary slide with the four header tables."""
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    # the title holds the audit title and, on a second line, the date
    cover.shapes.title.text_frame.text = "25123-0070 Claims handling (DRAFT)\x0b01/10/2026"
    contents, summary = (
        deck.slides.add_slide(deck.slide_layouts[6]) for _ in range(2)
    )
    icon = contents.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(1), Inches(1)
    )
    icon.click_action.target_slide = summary
    contents.shapes.add_textbox(
        Inches(1), Inches(2.2), Inches(3), Inches(0.5)
    ).text_frame.text = "Executive summary"
    for column, label in enumerate(
        ("Audit title", "Domain", "Process risk (gross)", "Key figures")
    ):
        table = summary.shapes.add_table(
            1, 2, Inches(0.5 + 2.2 * column), Inches(0.3), Inches(2), Inches(0.3)
        ).table
        table.cell(0, 0).text_frame.text = label
        table.cell(0, 1).text_frame.paragraphs[0].add_run().text = "XXX"
    return deck


def header_values(deck, slide):
    return {
        shape.table.cell(0, 0).text: shape.table.cell(0, 1).text
        for shape in list(deck.slides)[slide].shapes
        if getattr(shape, "has_table", False)
        and len(shape.table.rows) == 1
        and len(shape.table.columns) == 2
    }


def test_header_gets_title_choices_and_a_placeholder_for_key_figures():
    deck = make_header_deck()
    fill_header(deck, "Finance", "Material")

    assert header_values(deck, 2) == {
        "Audit title": "25123-0070 Claims handling (DRAFT)",  # without the date
        "Domain": "Finance",
        "Process risk (gross)": "Material",
        "Key figures": KEY_FIGURES_PLACEHOLDER,
    }


def test_header_values_that_are_not_given_keep_their_placeholder():
    deck = make_header_deck()
    fill_header(deck)

    values = header_values(deck, 2)
    assert values["Domain"] == "XXX" and values["Process risk (gross)"] == "XXX"
    assert values["Key figures"] == KEY_FIGURES_PLACEHOLDER


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not in checkout")
def test_template_header(tmp_path):
    out_file = tmp_path / "template.pptx"
    write_summary_to_deck(
        TEMPLATE, SUMMARY, out_file, domain="HR", process_risk="Major"
    )

    assert header_values(Presentation(str(out_file)), 3) == {
        "Audit title": "25XXX-0070 [AUDIT NAME] (DRAFT)",
        "Domain": "HR",
        "Process risk (gross)": "Major",
        "Key figures": KEY_FIGURES_PLACEHOLDER,
    }


# -------------------- Finding owner --------------------

FINDING_SLIDES = [
    ("0001 – Purchases are not authorised", "Risk: High",
     "We recommend the Procurement Manager (PUR-01) to:\nConfigure approval workflows."),
    ("0002 – Receipts are not documented", "Risk: Medium",
     "We recommend the Store Operations Manager (OPS-01) to:\nUse a receipt checklist."),
    ("0003 – Reviews are not documented", "Risk: Low", "Document the annual reviews."),
]  # fmt: skip


def make_owner_deck():
    """make_deck with one slide per finding: its title, its risk and its recommendation."""
    deck = make_deck()
    for title, risk, recommendation in FINDING_SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Findings and recommendations"
        table = slide.shapes.add_table(
            3, 3, Inches(1), Inches(2), Inches(8), Inches(2)
        ).table
        for column, text in enumerate((title, "", risk)):
            table.cell(0, column).text = text
        table.cell(1, 0).text = "Finding"
        table.cell(2, 0).text = "Recommendation"
        table.cell(2, 1).text = recommendation
    return deck


def test_owner_is_read_from_the_recommendation_of_each_finding():
    table = read_finding_table(make_owner_deck())

    assert table == [
        {"title": "0001 – Purchases are not authorised", "risk": "High",
         "owner": "Procurement Manager (PUR-01)"},
        {"title": "0002 – Receipts are not documented", "risk": "Medium",
         "owner": "Store Operations Manager (OPS-01)"},
        # a recommendation that does not start with "We recommend the ... to:" names no owner
        {"title": "0003 – Reviews are not documented", "risk": "Low", "owner": None},
    ]  # fmt: skip
    # a title is matched whatever its dash and case, or by its number
    assert finding_owner("0002 - receipts are not documented", table).startswith("Store")
    assert finding_owner("0001: Purchases without approval", table).startswith("Procurement")
    assert finding_owner("Something else", table) is None
    assert finding_owner("", table) is None


def test_owner_goes_in_the_third_column_of_the_findings_table():
    deck = make_owner_deck()
    findings = {
        "Receipts are not documented.": "Use a checklist.",
        "Reviews are not documented.": "Document them.",
        "A finding without a title.": "Fix it.",
    }
    titles = {
        "Receipts are not documented.": "0002 – Receipts are not documented",
        "Reviews are not documented.": "0003 – Reviews are not documented",
        "A finding without a title.": "",
    }
    fill_findings_table(deck, findings, "gpt-5.1", titles)

    owners = [row[2] for row in findings_rows(deck)]
    # no owner named on the slide, or no title: the cell stays for the audit manager
    assert owners == ["Store Operations Manager (OPS-01)", "", "", ""]


def test_owner_cell_keeps_its_placeholder_when_no_owner_is_found():
    deck = make_owner_deck()
    fill_findings_table(
        deck, {"Reviews are not documented.": "Document them."}, "gpt-5.1",
        {"Reviews are not documented.": "0003 – Reviews are not documented"},
    )  # fmt: skip

    assert findings_rows(deck)[0][2] == "[FINDING OWNER]"
