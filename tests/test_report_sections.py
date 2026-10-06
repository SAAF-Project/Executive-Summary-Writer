"""Tests for locating the sections of an audit report deck from its table of contents slide.

pytest tests/
"""

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from material import DEFAULT_SOURCES  # noqa: E402
from report_sections import (  # noqa: E402
    load_section,
    load_slide,
    locate_sections,
)

TEMPLATE = DEFAULT_SOURCES["pptx"]


def make_deck():
    """Cover, table of contents, then two sections of two and one slides."""
    deck = Presentation()
    blank = deck.slide_layouts[6]
    slides = [deck.slides.add_slide(blank) for _ in range(5)]
    for number, slide in enumerate(slides, start=1):
        box = slide.shapes.add_textbox(
            Inches(1), Inches(6), Inches(4), Inches(0.5)
        )
        box.text_frame.text = f"Body of slide {number}"

    contents = slides[1]
    # the second section comes first on the slide: the order must come from the links
    for column, (name, target) in enumerate(
        [("Findings", slides[4]), ("Introduction", slides[2])]
    ):
        left = Inches(1 + 4 * column)
        icon = contents.shapes.add_shape(
            MSO_SHAPE.OVAL, left, Inches(1), Inches(1), Inches(1)
        )
        icon.click_action.target_slide = target
        label = contents.shapes.add_textbox(
            left, Inches(2.2), Inches(2), Inches(0.5)
        )
        label.text_frame.text = name
    return deck


def test_sections_run_from_their_link_to_the_next_section():
    assert locate_sections(make_deck()) == {
        "Introduction": [3, 4],
        "Findings": [5],
    }


def test_section_text_leaves_out_the_navigation_bar():
    deck = make_deck()
    nav = list(deck.slides)[2].shapes.add_textbox(0, 0, Inches(2), Inches(0.3))
    nav.text_frame.text = "Findings"

    text = load_section(deck, "Introduction")
    assert "--- Slide 3 ---\nBody of slide 3" in text
    assert "--- Slide 4 ---\nBody of slide 4" in text
    assert "Findings" not in text
    assert "Findings" in load_slide(deck, 3)  # the plain slide keeps everything


def test_unknown_section_lists_the_sections():
    with pytest.raises(ValueError, match="The deck has: Introduction, Findings"):
        load_section(make_deck(), "Appendix")


def test_deck_without_links_is_refused():
    deck = Presentation()
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6])
    with pytest.raises(ValueError, match="no links to sections"):
        locate_sections(deck)


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not in checkout")
def test_template():
    deck = Presentation(str(TEMPLATE))
    assert locate_sections(deck) == {
        "Introduction": [3],
        "Executive summary": [4, 5],
        "Process dashboard": [6],
        "Finding overview": [7],
        "Findings and recommendations": [8, 9, 10, 11, 12, 13, 14],
        "Appendices": [15, 16, 17],
    }
    assert "[FINDING DESCRIPTION + SHORT CONTEXT]" in load_section(
        deck, "Findings and recommendations"
    )
