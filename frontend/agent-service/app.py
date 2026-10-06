"""Local hackathon HTTP adapter for the original Executive Summary Writer graph."""

import copy
import hmac
import io
import json
import os
import sys
import threading
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
checkout = os.getenv("AGENT_REPOSITORY_PATH")
agent_path = Path(checkout).expanduser().resolve() / "scripts" if checkout else ROOT / "upstream"
if not (agent_path / "graph.py").is_file():
    raise RuntimeError("AGENT_REPOSITORY_PATH must point to the Executive-Summary-Writer checkout.")
sys.path.insert(0, str(agent_path))

import anthropic  # noqa: E402
import prompts  # noqa: E402
from graph import build_graph  # noqa: E402
from langgraph.types import Command  # noqa: E402
from llm import ClaudeLLM, MODEL_NAME  # noqa: E402
from fastapi import FastAPI, HTTPException, Request  # noqa: E402
from fastapi.responses import JSONResponse  # noqa: E402
from fastapi.exceptions import RequestValidationError  # noqa: E402
from pydantic import BaseModel, Field  # noqa: E402
from presentation import PresentationClaude, InsufficientGradeEvidence, deck_facts, extract_presentation, draft_fields  # noqa: E402
from typing import Literal  # noqa: E402

ClaudeLLM = PresentationClaude

REVISION = "cd950925efacd756c6d9da5a10ce7b502b3b0ac3"
MAX_BODY = 27 * 1024 * 1024
MAX_FILE = 4 * 1024 * 1024
MAX_SESSIONS = 30
SESSION_TTL = 2 * 60 * 60


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def message(role: str, content: str) -> dict:
    return {"id": str(uuid.uuid4()), "role": role, "content": content, "createdAt": now()}


class LazyClaude:
    """The first graph question needs no model call; initialize Claude when needed."""

    def __init__(self):
        self.provider = None
        self.api_key = None

    def __getattr__(self, key):
        api_key = os.getenv("ANTHROPIC_API_KEY")
        if self.provider is None or self.api_key != api_key:
            self.provider = ClaudeLLM()
            self.api_key = api_key
        return getattr(self.provider, key)


@dataclass
class Session:
    id: str
    graph: Any
    template: dict
    provider: Any = None
    material: list = field(default_factory=list)
    plan: dict | None = None
    analysis: dict | None = None
    initial_input: Any = None
    result: dict | None = None
    phase: str = "interview"
    needs_grade_evidence: bool = False
    pending_material_update: bool = False
    revision: int = 0
    status: str = "processing"
    current_step: str | None = None
    options: list = field(default_factory=list)  # fixed answers to the current question, for a selection menu
    messages: list = field(default_factory=list)
    confirmed: dict = field(default_factory=dict)
    artifacts: list = field(default_factory=list)
    error: dict | None = None
    seen: set = field(default_factory=set)
    touched: float = field(default_factory=time.monotonic)
    lock: Any = field(default_factory=threading.Lock)

    @property
    def config(self):
        return {"configurable": {"thread_id": self.id}}


class AgentService:
    def __init__(self, provider_factory: Callable = LazyClaude):
        self.provider_factory = provider_factory
        self.sessions: dict[str, Session] = {}
        self.lock = threading.Lock()
        self.executor = ThreadPoolExecutor(max_workers=4)

    def get(self, session_id: str) -> Session:
        with self.lock:
            session = self.sessions.get(session_id)
        if not session:
            raise HTTPException(404, "This conversation ended when the local agent restarted. Start a new conversation.")
        session.touched = time.monotonic()
        return session

    def add(self, template: dict, material=None, plan=None) -> Session:
        with self.lock:
            expired = [key for key, value in self.sessions.items() if value.status != "processing" and time.monotonic() - value.touched > SESSION_TTL]
            for key in expired:
                del self.sessions[key]
            if len(self.sessions) >= MAX_SESSIONS:
                raise HTTPException(429, "The local demo has reached its conversation limit. Restart the agent to clear old sessions.")
            provider = self.provider_factory()
            session = Session(str(uuid.uuid4()), build_graph(provider), template, provider=provider, material=material or [], plan=plan, phase="analysis" if plan else "interview")
            self.sessions[session.id] = session
            return session

    def run(self, session: Session, invocation: Any):
        try:
            if session.pending_material_update:
                # Re-enter the original grading node with the clarification, without repeating earlier answers.
                session.graph.update_state(session.config, {"material": session.material}, as_node=f"{prompts.POSITIVES.key}_interpret")
                session.pending_material_update = False
            if session.plan and session.analysis is None:
                analysis = session.provider.analyze_presentation(session.material)
                with session.lock:
                    session.analysis = analysis.model_dump()
                    session.phase = "interview"
                invocation = {**session.initial_input, "proposals": {prompts.STEPS[0].key: analysis.firstQuestion}}
            elif session.plan and session.result is None and invocation is None and not session.graph.get_state(session.config).values:
                invocation = {**session.initial_input, "proposals": {prompts.STEPS[0].key: session.analysis["firstQuestion"]}}
            result = session.result if session.phase == "format" and session.result else session.graph.invoke(invocation, session.config)
            if session.plan and not result.get("__interrupt__"):
                session.result = result
                session.phase = "format"
                fields, draft_warnings = draft_fields(result, session.plan, MODEL_NAME)
            with session.lock:
                session.confirmed = {step.key: result.get(step.key) for step in prompts.STEPS if step.key in result.get("done", [])}
                for key in ("grade", "tone", "domain", "process_risk"):
                    if result.get(key): session.confirmed[key] = result[key]
                if session.analysis and result.get("grade_suggestion"):
                    references = sorted({number for number in getattr(session.provider, "grade_reference_slides", []) if isinstance(number, int) and 1 <= number <= len(session.template["slides"])})
                    session.analysis["processRisk"].update({"grade": result["grade_suggestion"], "rationale": result.get("grade_reason", ""), "evidenceSlides": references})
                interrupts = result.get("__interrupt__", ())
                if interrupts:
                    question = interrupts[0].value
                    session.current_step = question["step"]
                    session.options = list(question.get("options", []))
                    session.messages.append(message("assistant", question["message"]))
                    session.status = "awaiting-input"
                else:
                    session.current_step = None
                    session.options = []
                    session.status = "awaiting-review"
                    session.artifacts = [{"id": str(uuid.uuid4()), "kind": "content", "title": "Executive summary draft", "mimeType": "text/markdown", "reviewStatus": "draft", "content": result.get("summary", ""), "confirmation": result.get("confirmation", "")}]
                    if session.plan:
                        session.artifacts[0].update({"fields": fields, "grade": result.get("grade"), "processRisk": session.analysis["processRisk"], "warnings": draft_warnings})
                    session.messages.append(message("assistant", "Your executive summary and process grade are ready to review in the presentation." if session.plan else "Your executive summary draft is ready. Review and edit it before using it in your report."))
                session.error = None
                session.revision += 1
                session.touched = time.monotonic()
        except InsufficientGradeEvidence as exc:
            checkpoint = session.graph.get_state(session.config).values
            with session.lock:
                session.confirmed = {step.key: checkpoint.get(step.key) for step in prompts.STEPS if step.key in checkpoint.get("done", [])}
                if session.analysis:
                    references = sorted({number for number in exc.references if isinstance(number, int) and 1 <= number <= len(session.template["slides"])})
                    session.analysis["processRisk"].update({"grade": None, "rationale": exc.reason, "evidenceSlides": references})
                session.needs_grade_evidence = True
                session.current_step = "grade"
                session.options = []
                session.status = "awaiting-input"
                session.messages.append(message("assistant", "More evidence is needed before recommending a process grade.\n\n" + exc.question))
                session.error = None
                session.revision += 1
                session.touched = time.monotonic()
        except Exception as exc:
            if isinstance(exc, anthropic.AuthenticationError):
                detail = "Claude did not accept the API key. Reconnect Claude with a valid key, then retry."
            elif isinstance(exc, anthropic.RateLimitError):
                detail = "Claude is rate limited. Wait a moment, then retry."
            elif isinstance(exc, (anthropic.APIConnectionError, anthropic.APITimeoutError)):
                detail = "The connection to Claude timed out. Check your connection, then retry."
            elif isinstance(exc, anthropic.BadRequestError):
                detail = "Claude could not process this request. Check the model setting and uploaded material, then retry."
            elif not os.getenv("ANTHROPIC_API_KEY"):
                detail = "Connect Claude with your Anthropic API key to continue this conversation."
            else:
                detail = "The agent could not complete this step. Retry, or check the local agent configuration."
            # Never expose SDK request details, key values, or source material in errors.
            with session.lock:
                session.status = "failed"
                session.error = {"code": "AGENT_STEP_FAILED", "message": detail, "retryable": True}
                session.revision += 1

    def queue(self, session: Session, invocation: Any):
        self.executor.submit(self.run, session, invocation)

    def snapshot(self, session: Session) -> dict:
        with session.lock:
            definitions = ([("analysis", "Read presentation")] if session.plan else []) + [(step.key, step.label) for step in prompts.STEPS] + [("grade", "Overall grade"), ("tone", "Tone"), ("domain", "Domain"), ("process_risk", "Process risk (gross)"), ("challenge", "Challenge review"), ("summary", "Presentation review" if session.plan else "Draft review")]
            current_index = next((index for index, (key, _) in enumerate(definitions) if key == session.current_step), len(definitions) - 1 if session.artifacts else -1)
            steps = []
            for index, (key, label) in enumerate(definitions):
                status = "pending"
                if key in session.confirmed:
                    status = "skipped" if session.confirmed[key] is None else "confirmed"
                elif index < current_index:
                    status = "confirmed"
                elif index == current_index or key == "analysis" and session.phase == "analysis":
                    status = "current"
                steps.append({"id": key, "label": label, "status": status})
            return copy.deepcopy({"contractVersion": 1, "sessionId": session.id, "revision": session.revision, "status": session.status, "phase": session.phase, "analysis": session.analysis, "steps": steps, "message": session.messages[-1]["content"] if session.messages and session.messages[-1]["role"] == "assistant" else None, "options": session.options if session.status == "awaiting-input" else [], "messages": session.messages, "confirmed": session.confirmed, "artifacts": session.artifacts, "error": session.error})


class Reply(BaseModel):
    expectedRevision: int = Field(ge=0)
    idempotencyKey: str = Field(min_length=8, max_length=100)
    reply: str = Field(min_length=1, max_length=20000)


class Retry(BaseModel):
    expectedRevision: int = Field(ge=0)


class Configuration(BaseModel):
    apiKey: str = Field(min_length=20, max_length=500)


class Review(BaseModel):
    expectedRevision: int = Field(ge=0)
    content: str = Field(min_length=1, max_length=50000)
    fields: dict[str, str] | None = None
    grade: Literal["A", "B", "C", "D"] | None = None


def read_material(name: str, data: bytes) -> dict:
    suffix = Path(name).suffix.lower()
    safe_name = Path(name).name[:160]
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise HTTPException(400, f"{safe_name} is not a PDF file.")
        import base64
        return {"type": "document", "title": safe_name, "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode("ascii")}}
    if suffix == ".docx":
        from docx import Document
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(400, f"{safe_name} expands beyond the demo's document limit.")
            document = Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            text = "\n".join(parts)
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(400, f"{safe_name} could not be read. Use a fresh .docx or text file.") from None
    elif suffix in {".txt", ".md"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(400, f"Save {safe_name} as UTF-8 text and try again.") from None
    else:
        raise HTTPException(400, "Audit material must be .txt, .md, .docx, or .pdf. Upload presentation layouts in Templates.")
    if not text.strip():
        raise HTTPException(400, f"{safe_name} has no readable text.")
    return {"type": "text", "text": f"Audit material from {safe_name}:\n{text.strip()}"}


def create_app(provider_factory: Callable = LazyClaude) -> FastAPI:
    if os.getenv("AGENT_HOSTED") == "1" and not os.getenv("AGENT_API_TOKEN"):
        raise RuntimeError("Set AGENT_API_TOKEN before starting a hosted agent.")
    app = FastAPI(title="Local audit report agent", docs_url=None, redoc_url=None)
    service = AgentService(provider_factory)
    app.state.service = service

    @app.exception_handler(RequestValidationError)
    async def validation_error(_request, _exception):
        return JSONResponse({"detail": "Check the required fields and try again."}, status_code=422)

    @app.middleware("http")
    async def local_adapter_header(request: Request, call_next):
        token = os.getenv("AGENT_API_TOKEN")
        if request.url.path != "/health" and token and not hmac.compare_digest(request.headers.get("authorization", "").encode(), f"Bearer {token}".encode()):
            return JSONResponse({"detail": "The agent service credentials are missing or invalid."}, status_code=401)
        if request.url.path != "/health" and request.headers.get("x-template-studio") != "local-demo":
            return JSONResponse({"detail": "Use the local ESWriter frontend."}, status_code=403)
        try:
            length = int(request.headers.get("content-length", "0"))
        except ValueError:
            return JSONResponse({"detail": "Invalid request size."}, status_code=400)
        if length > MAX_BODY:
            return JSONResponse({"detail": "The audit material is too large for this demo."}, status_code=413)
        return await call_next(request)

    @app.get("/health")
    def health():
        return {"status": "ok"}

    @app.get("/status")
    def status():
        return {"connected": True, "ready": bool(os.getenv("ANTHROPIC_API_KEY")), "canConfigure": not bool(os.getenv("AGENT_API_TOKEN")), "provider": "Claude / Anthropic", "model": MODEL_NAME, "contractVersion": 1, "acceptedMaterialTypes": [".pptx", ".txt", ".md", ".docx", ".pdf"], "outputs": ["content", "presentation"], "resumable": True, "supportsTemplateBindings": True, "sessionPersistence": "local-process", "repository": {"url": "https://github.com/SAAF-Project/Executive-Summary-Writer", "revision": "local-checkout" if checkout else REVISION}}

    @app.post("/configuration")
    def configure(body: Configuration):
        if os.getenv("AGENT_API_TOKEN"):
            raise HTTPException(403, "Configure Claude in the hosted agent's environment settings.")
        if any(session.status == "processing" for session in service.sessions.values()):
            raise HTTPException(409, "Wait for the current agent step to finish before reconnecting Claude.")
        try:
            with anthropic.Anthropic(api_key=body.apiKey.strip(), timeout=15, max_retries=0) as client:
                client.models.list(limit=1)
        except anthropic.AuthenticationError:
            raise HTTPException(401, "Claude did not accept that API key. Check the key and try again.") from None
        except Exception:
            raise HTTPException(502, "Could not connect to Claude. Check your connection and key permissions, then try again.") from None
        os.environ["ANTHROPIC_API_KEY"] = body.apiKey.strip()
        return status()

    @app.post("/sessions", status_code=202)
    async def start(request: Request):
        form = await request.form(max_files=6, max_fields=5, max_part_size=MAX_BODY)
        try:
            template = json.loads(str(form.get("template", "{}")))
        except json.JSONDecodeError:
            raise HTTPException(400, "The template definition could not be read.") from None
        if not isinstance(template, dict) or not isinstance(template.get("name"), str) or not isinstance(template.get("slides"), list) or not 1 <= len(template["slides"]) <= 100:
            raise HTTPException(400, "Choose a saved presentation template.")
        notes = str(form.get("notes", "")).strip()
        request_text = str(form.get("request", "")).strip()
        if len(notes) > 100000 or not 1 <= len(request_text) <= 4000:
            raise HTTPException(400, "Add a report request, up to 4,000 characters, and keep audit notes below 100,000 characters.")
        material = [{"type": "text", "text": f"Audit manager's notes:\n{notes}"}] if notes else []
        plan = None
        presentation = form.get("presentation")
        if presentation is not None:
            if not hasattr(presentation, "read") or not (presentation.filename or "").lower().endswith(".pptx"):
                raise HTTPException(400, "Upload an audit PowerPoint in .pptx format.")
            data = await presentation.read(25 * 1024 * 1024 + 1)
            await presentation.close()
            try:
                evidence, valid_ids, slide_count = extract_presentation(data)
                plan = json.loads(str(form.get("plan", "{}")))
                targets = plan.get("summaryFields", [])
                if not 1 <= len(targets) <= 50 or len({t["id"] for t in targets}) != len(targets) or any(t["id"] not in valid_ids or not isinstance(t["label"], str) or not 1 <= t["maxChars"] <= 6000 for t in targets):
                    raise ValueError("The executive-summary areas could not be located. Check the presentation layout.")
                if slide_count != len(template["slides"]):
                    raise ValueError("The presentation and selected report do not match.")
                plan.update(deck_facts(data))
            except (ValueError, KeyError, TypeError, zipfile.BadZipFile, ET.ParseError) as exc:
                raise HTTPException(400, "The presentation or its executive-summary mapping could not be read. Upload a standard .pptx with a labeled executive-summary slide.") from None
            material = evidence + material
            del data
        total = 0
        for upload in form.getlist("files"):
            if not hasattr(upload, "read"):
                raise HTTPException(400, "Invalid audit attachment.")
            data = await upload.read(MAX_FILE + 1)
            total += len(data)
            if not data or len(data) > MAX_FILE or total > 12 * 1024 * 1024:
                raise HTTPException(400, "Use up to five nonempty audit files, 4 MB each and 12 MB in total.")
            material.append(read_material(upload.filename or "material", data))
            await upload.close()
        if not material:
            raise HTTPException(400, "Add audit notes or attach audit source material. A presentation layout does not supply audit evidence.")
        session = service.add(template, material, plan)
        session.messages.append(message("user", request_text))
        session.initial_input = {"material": material, "request": request_text}
        service.queue(session, session.initial_input)
        return service.snapshot(session)

    @app.get("/sessions/{session_id}")
    def get_session(session_id: str):
        return service.snapshot(service.get(session_id))

    @app.post("/sessions/{session_id}/reply", status_code=202)
    def reply(session_id: str, body: Reply):
        session = service.get(session_id)
        with session.lock:
            if body.idempotencyKey in session.seen:
                duplicate = True
            else:
                duplicate = False
                if session.revision != body.expectedRevision or session.status != "awaiting-input":
                    raise HTTPException(409, "The conversation has moved on. Refresh its latest state before replying.")
                if not body.reply.strip():
                    raise HTTPException(400, "Enter a reply to continue.")
                if session.needs_grade_evidence and sum(len(item.get("text", "")) for item in session.material) + len(body.reply) > 200000:
                    raise HTTPException(400, "Use a shorter evidence clarification for this report.")
                session.seen.add(body.idempotencyKey)
                session.messages.append(message("user", body.reply.strip()))
                if session.needs_grade_evidence:
                    session.material.append({"type": "text", "text": "Additional audit evidence from the audit manager:\n" + body.reply.strip()})
                    session.pending_material_update = True
                    session.needs_grade_evidence = False
                session.status = "processing"
                session.revision += 1
        if not duplicate:
            service.queue(session, None if session.pending_material_update else Command(resume=body.reply.strip()))
        return service.snapshot(session)

    @app.post("/sessions/{session_id}/retry", status_code=202)
    def retry(session_id: str, body: Retry):
        session = service.get(session_id)
        with session.lock:
            if session.revision != body.expectedRevision or session.status != "failed":
                raise HTTPException(409, "The conversation has moved on. Refresh its latest state before retrying.")
            session.status = "processing"
            session.error = None
            session.revision += 1
        service.queue(session, None)
        return service.snapshot(session)

    @app.post("/sessions/{session_id}/review")
    def review(session_id: str, body: Review):
        session = service.get(session_id)
        with session.lock:
            if session.revision != body.expectedRevision or session.status not in {"awaiting-review", "completed"} or not session.artifacts:
                raise HTTPException(409, "Refresh the draft before approving it.")
            if not body.content.strip():
                raise HTTPException(400, "The reviewed draft cannot be empty.")
            if session.plan:
                allowed = {target["id"] for target in session.plan["summaryFields"]}
                if body.fields is None or set(body.fields) != allowed or any(len(value) > next(t["maxChars"] for t in session.plan["summaryFields"] if t["id"] == key) for key, value in body.fields.items()):
                    raise HTTPException(400, "Review all summary cells and shorten any text that exceeds its slide limit.")
                session.artifacts[0]["fields"] = body.fields
                session.artifacts[0]["reviewedGrade"] = body.grade
            session.artifacts[0]["content"] = body.content.strip()
            session.artifacts[0]["reviewStatus"] = "approved"
            session.status = "completed"
            session.revision += 1
            session.messages.append(message("assistant", "Your reviewed presentation is ready to download. ESWriter exports your summary edits and selected process grade into a copy of the original PowerPoint." if session.plan else "Your reviewed draft is ready to download. The source presentation is unchanged."))
        return service.snapshot(session)

    return app


app = create_app()
