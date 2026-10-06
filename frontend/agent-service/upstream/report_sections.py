"""Finds where each part of an audit report deck is, using its table of contents slide.

The report template (samples/Template - 0070 and 0010 - Audit report Anonymized.pptx) has a fixed
table of contents on slide 2: one icon per section that links to the slide where the section starts,
with the section name below it. `locate_sections` reads those links into a dictionary, so the rest of
the code can ask for a section by name instead of by slide number.

    python scripts/report_sections.py            # print the dictionary for the template
    python scripts/report_sections.py report.pptx
"""

import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from material import (
    DEFAULT_SOURCES,
    FINDING_SLIDE_TITLE,
    _shape_text,
    _slide_title,
)

CONTENTS_SLIDE = 2
# the risk of a finding is fixed by its author as "Risk: High", "Risk: Medium" or "Risk: Low"
RISK_LEVELS = ("High", "Medium", "Low")


def _centre(shape: Any) -> Tuple[float, float]:
    return shape.left + shape.width / 2, shape.top + shape.height / 2


def _linked_slide(shape: Any) -> Optional[Any]:
    """The slide a shape jumps to when clicked, if any."""
    try:
        return shape.click_action.target_slide
    except (AttributeError, KeyError):
        return None


def _label(shape: Any) -> str:
    return shape.text_frame.text.strip() if shape.has_text_frame else ""


def locate_sections(presentation: Any) -> Dict[str, List[int]]:
    """Section name -> the slide numbers (1-based) that belong to it, in the order of the deck.

    A section starts at the slide its icon on the table of contents links to and runs until the next
    section starts. For the template:

        {"Introduction": [3], "Executive summary": [4, 5], "Process dashboard": [6],
         "Finding overview": [7], "Findings and recommendations": [8, ..., 14], "Appendices": [15, 16, 17]}
    """
    slides = list(presentation.slides)
    if len(slides) < CONTENTS_SLIDE:
        raise ValueError(
            f"The deck has no slide {CONTENTS_SLIDE} for the table of contents."
        )
    number_of = {slide.slide_id: n for n, slide in enumerate(slides, start=1)}
    shapes = list(slides[CONTENTS_SLIDE - 1].shapes)
    labels = [s for s in shapes if _label(s) and _linked_slide(s) is None]

    starts: Dict[str, int] = {}
    for icon in shapes:
        target = _linked_slide(icon)
        if target is None:
            continue
        # the name of a section is the text closest below its icon
        x, y = _centre(icon)
        below = [s for s in labels if s.top >= icon.top]
        if not below:
            continue
        name = min(
            below,
            key=lambda s: (_centre(s)[0] - x) ** 2 + (_centre(s)[1] - y) ** 2,
        )
        starts[_label(name)] = number_of[target.slide_id]
    if not starts:
        raise ValueError(
            f"Slide {CONTENTS_SLIDE} has no links to sections. Is this the audit report template?"
        )

    ordered = sorted(starts.items(), key=lambda item: item[1])
    ends = [start for _, start in ordered[1:]] + [len(slides) + 1]
    return {
        name: list(range(start, end))
        for (name, start), end in zip(ordered, ends)
    }


def load_slide(
    presentation: Any, number: int, sections: Optional[Dict[str, List[int]]] = None
) -> str:
    """The text of one slide. With `sections`, the navigation bar and the slide number are left out."""
    slide = list(presentation.slides)[number - 1]
    skip = set(sections or ()) | ({str(number)} if sections else set())
    lines = []
    for shape in slide.shapes:
        if _label(shape) in skip:
            continue
        lines.extend(_shape_text(shape))
    return "\n".join(lines)


def load_section(presentation: Any, name: str) -> str:
    """The text of every slide of one section, by the name it has on the table of contents."""
    sections = locate_sections(presentation)
    if name not in sections:
        raise ValueError(
            f"No section '{name}'. The deck has: {', '.join(sections)}"
        )
    return "\n\n".join(
        f"--- Slide {number} ---\n{load_slide(presentation, number, sections)}"
        for number in sections[name]
    )


def read_findings(presentation: Any) -> List[Tuple[str, Optional[str]]]:
    """(finding title, risk level) for every slide with the title "Findings and recommendations".

    Both come from the first row of the table on the slide: the title of the finding, and next to it
    "Risk: High", "Risk: Medium" or "Risk: Low". The level is None when it is none of these, as in the
    empty template ("Risk: xxx").
    """
    findings: List[Tuple[str, Optional[str]]] = []
    for slide in presentation.slides:
        if _slide_title(slide) != FINDING_SLIDE_TITLE:
            continue
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            title, *details = [
                " ".join(cell.text.split())
                for cell in list(shape.table.rows)[0].cells
                if not cell.is_spanned
            ]
            stated = re.search(r"Risk:\s*(\w+)", " ".join(details), re.IGNORECASE)
            level = stated.group(1).capitalize() if stated else None
            findings.append((title, level if level in RISK_LEVELS else None))
    return findings


def count_risks(findings: List[Tuple[str, Optional[str]]]) -> List[int]:
    """The number of findings per risk level, in the order High, Medium, Low."""
    return [
        sum(1 for _, level in findings if level == wanted)
        for wanted in RISK_LEVELS
    ]


if __name__ == "__main__":
    from pptx import Presentation

    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCES["pptx"]
    for section, numbers in locate_sections(Presentation(str(path))).items():
        print(f"{section}: slides {', '.join(map(str, numbers))}")
    for title, level in read_findings(Presentation(str(path))):
        print(f"Finding: {title} (risk: {level or 'not recognised'})")

