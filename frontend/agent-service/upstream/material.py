"""Loads the audit material (draft report, observations, management responses) as Claude content blocks.

The input source type says how a file is read. It is given explicitly (`--input-source-type`, or by
the interface that uploads the file) or, when left out, taken from the file extension.
"""

import base64
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

SOURCE_TYPES = ("txt", "md", "doc", "docx", "xls", "xlsx", "ppt", "pptx", "pdf")
DEFAULT_SOURCE_TYPE = "pptx"

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
# loaded when no file is given
DEFAULT_SOURCES = {
    "pptx": SAMPLES / "Template - 0070 and 0010 - Audit report Anonymized.pptx",
    "md": SAMPLES / "synthetic-audit-report.md",
}

# a slide with this title holds one finding, as a table of labels and contents
FINDING_SLIDE_TITLE = "Findings and recommendations"

# old binary Office formats are converted to the current format first
LEGACY_TYPES = {"doc": "docx", "xls": "xlsx", "ppt": "pptx"}


def _row(cells: Iterable[Any]) -> str:
    return " | ".join("" if c is None else str(c).strip() for c in cells)


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(_row(cell.text for cell in row.cells))
    return "\n".join(parts)


def _shape_text(shape: Any) -> List[str]:
    """Text of one slide shape; grouped shapes are read one by one."""
    if hasattr(shape, "shapes"):
        return [line for child in shape.shapes for line in _shape_text(child)]
    if getattr(shape, "has_table", False):
        rows = [_row(cell.text for cell in row.cells) for row in shape.table.rows]
        return [row for row in rows if row.replace("|", "").strip()]
    if shape.has_text_frame and shape.text_frame.text.strip():
        return [shape.text_frame.text.strip()]
    return []


def _slide_title(slide: Any) -> str:
    title = slide.shapes.title
    return title.text_frame.text.strip() if title is not None else ""


def _finding_text(table: Any) -> List[str]:
    """The table of a finding slide as plain text.

    The first row holds the title of the finding. In every row below it, the first column is a fixed
    label (Finding, Root Cause(s), Risk(s), Stake(s), Recommendation) and the second column its content:

        Finding title: 0123 - Access rights are not reviewed
        Risk: High
        Finding: ...
        Root Cause(s): ...
        Recommendation: ...
    """
    # a merged cell keeps hidden text in the cells it covers, so those are left out
    rows = [
        [cell.text.strip() for cell in row.cells if not cell.is_spanned]
        for row in table.rows
    ]
    (title, *details), *body = rows
    lines = [f"Finding title: {title}"] + [d for d in details if d]
    for label, *cells in body:
        content, *extra = [c for c in cells if c] or [""]
        if not label or not content:
            continue
        note = f" ({', '.join(extra)})" if extra else ""
        lines.append(f"{label}: {content}{note}")
    return lines


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    parts = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        is_finding = _slide_title(slide) == FINDING_SLIDE_TITLE
        lines = [
            line
            for shape in slide.shapes
            for line in (
                _finding_text(shape.table)
                if is_finding and getattr(shape, "has_table", False)
                else _shape_text(shape)
            )
        ]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        if lines:
            parts.append(f"--- Slide {number} ---\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    # data_only: the calculated values, not the formulas
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows = [_row(row) for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if row.replace("|", "").strip()]
        if rows:
            parts.append(f"--- Sheet {sheet.title} ---\n" + "\n".join(rows))
    workbook.close()
    return "\n\n".join(parts)


READERS = {"docx": _docx_text, "xlsx": _xlsx_text, "pptx": _pptx_text}


def _legacy_text(path: Path, source_type: str) -> str:
    """Convert a .doc / .xls / .ppt file with LibreOffice and read the result."""
    target = LEGACY_TYPES[source_type]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError(
            f"Reading .{source_type} files needs LibreOffice, which was not found. "
            f"Save {path.name} as .{target} and use that instead."
        )
    with tempfile.TemporaryDirectory() as tmp:
        # the extension tells LibreOffice which format it is reading
        source = Path(tmp) / f"source.{source_type}"
        shutil.copyfile(path, source)
        subprocess.run(
            [soffice, "--headless", "--convert-to", target, "--outdir", tmp, str(source)],
            check=True,
            capture_output=True,
            timeout=120,
        )
        return READERS[target](Path(tmp) / f"source.{target}")


def default_source(source_type: str) -> Path:
    """The sample that is loaded when no file is given."""
    if source_type not in DEFAULT_SOURCES:
        raise ValueError(
            f"There is no sample of type {source_type}. Give the file to read."
        )
    return DEFAULT_SOURCES[source_type]


def load_material(
    paths: List[Path], source_type: Optional[str] = None
) -> List[Dict[str, Any]]:
    """One content block per file: a document block for pdf, text for every other type.

    `source_type` is one of SOURCE_TYPES and applies to all files. Without it, each file is read
    according to its extension.
    """
    blocks: List[Dict[str, Any]] = []
    for path in paths:
        kind = source_type or path.suffix.lower().lstrip(".")
        if kind not in SOURCE_TYPES:
            raise ValueError(
                f"Unsupported file type: {path.name} (use {', '.join(SOURCE_TYPES)})"
            )
        if kind == "pdf":
            data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
            blocks.append(
                {
                    "type": "document",
                    "title": path.name,
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": data,
                    },
                }
            )
            continue
        try:
            if kind in ("txt", "md"):
                text = path.read_text(encoding="utf-8", errors="ignore")
            elif kind in LEGACY_TYPES:
                text = _legacy_text(path, kind)
            else:
                text = READERS[kind](path)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(
                f"Could not read {path.name} as .{kind}: {exc}"
            ) from exc
        if not text.strip():
            raise ValueError(f"No text found in {path.name}")
        blocks.append(
            {
                "type": "text",
                "text": f'<audit_material file="{path.name}">\n{text.strip()}\n</audit_material>',
            }
        )
    return blocks

