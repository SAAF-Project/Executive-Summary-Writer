"""Adapter tests use explicit stand-ins; the runnable demo has no fake provider."""

import json
import sys
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import create_app
from llm import Interpretation, BoardSummary, GradeSuggestion, Finding

HEADERS = {"X-Template-Studio": "local-demo"}
TEMPLATE = {"name": "Test presentation layout", "slides": [{"title": "Executive Summary"}]}


class ScriptedModel:
    def interpret(self, question, reply, proposal):
        return Interpretation(decision="skipped" if reply == "skip" else "confirmed", text=reply)

    def refine_storyline(self, storyline, material, confirmed):
        return storyline

    def interpret_tone(self, reply):
        return "Balanced"

    def challenge(self, material, confirmed):
        return ["Who owns the action plan?"]

    def suggest_grade(self, material):
        return GradeSuggestion(grade="B", reason="Explicit test-only grade rationale.")

    def summarise_findings(self, material):
        return [Finding(finding="Explicit test finding", recommendation="Explicit test recommendation")]

    def write_summary(self, material, confirmed, tone, request, questions, answers):
        assert material[0]["text"].endswith("Synthetic audit notes for a test.")
        return BoardSummary(exe_summary="## Executive Board Summary\nExplicit test-fixture draft.", pos_points=["Explicit test-positive point."])


@pytest.fixture
def client(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = create_app(ScriptedModel)
    with TestClient(app, headers=HEADERS) as client:
        yield client
    app.state.service.executor.shutdown(wait=True)


def start(client):
    response = client.post("/sessions", data={"template": json.dumps(TEMPLATE), "request": "Write a concise executive summary.", "notes": "Synthetic audit notes for a test."})
    assert response.status_code == 202, response.text
    return settled(client, response.json()["sessionId"])


def settled(client, session_id):
    for _ in range(100):
        response = client.get(f"/sessions/{session_id}")
        assert response.status_code == 200
        body = response.json()
        if body["status"] != "processing":
            return body
        time.sleep(0.02)
    raise AssertionError("Agent adapter job did not finish.")


def reply(client, session, text, key):
    response = client.post(f"/sessions/{session['sessionId']}/reply", json={"expectedRevision": session["revision"], "idempotencyKey": key, "reply": text})
    assert response.status_code == 202, response.text
    return settled(client, session["sessionId"])


def test_original_graph_questions_replies_review_and_idempotency(client):
    session = start(client)
    assert session["steps"][0]["status"] == "current"
    assert session["status"] == "awaiting-input"
    assert "root cause" in session["message"].lower()
    old_revision = session["revision"]
    session = reply(client, session, "Ownership is unclear.", "root-reply-unique")
    repeated = client.post(f"/sessions/{session['sessionId']}/reply", json={"expectedRevision": old_revision, "idempotencyKey": "root-reply-unique", "reply": "Ownership is unclear."})
    assert repeated.status_code == 202
    assert len(repeated.json()["messages"]) == len(session["messages"])
    assert session["confirmed"]["root_cause"] == "Ownership is unclear."
    session = reply(client, session, "skip", "relationships-reply")
    assert session["confirmed"]["relationships"] is None
    session = reply(client, session, "Actions need accountable owners.", "storyline-reply")
    session = reply(client, session, "Controls are documented.", "positive-reply-unique")
    assert session["steps"][4]["status"] == "current"
    invalid_grade = reply(client, session, "I have not decided", "grade-invalid-unique")
    assert invalid_grade["steps"][4]["status"] == "current"
    session = reply(client, invalid_grade, "C", "grade-reply-unique")
    assert session["confirmed"]["grade"] == "C"
    session = reply(client, session, "Balanced", "tone-reply-unique")
    assert session["steps"][6]["status"] == "current"
    session = reply(client, session, "The process owner.", "challenge-reply-unique")
    assert session["status"] == "awaiting-review"
    assert session["artifacts"][0]["reviewStatus"] == "draft"
    response = client.post(f"/sessions/{session['sessionId']}/review", json={"expectedRevision": session["revision"], "content": "Edited and reviewed test draft."})
    assert response.status_code == 200
    assert response.json()["status"] == "completed"
    assert response.json()["artifacts"][0]["content"] == "Edited and reviewed test draft."
    assert response.json()["artifacts"][0]["reviewStatus"] == "approved"


def test_stale_replies_and_invalid_inputs_are_rejected(client):
    session = start(client)
    stale = client.post(f"/sessions/{session['sessionId']}/reply", json={"expectedRevision": 999, "idempotencyKey": "stale-reply-unique", "reply": "Stale answer"})
    assert stale.status_code == 409
    invalid = client.post("/sessions", data={"template": json.dumps(TEMPLATE), "request": "Write report"})
    assert invalid.status_code == 400
    wrong_file = client.post("/sessions", data={"template": json.dumps(TEMPLATE), "request": "Write report"}, files={"files": ("layout.pptx", b"test")})
    assert wrong_file.status_code == 400
    secret_fixture = "not-a-real-key"
    configuration = client.post("/configuration", json={"apiKey": secret_fixture})
    assert configuration.status_code == 422
    assert secret_fixture not in configuration.text
    assert client.get("/sessions/00000000-0000-0000-0000-000000000000").status_code == 404


def test_browser_cannot_call_python_directly_without_adapter_header():
    app = create_app(ScriptedModel)
    with TestClient(app) as client:
        assert client.get("/health").status_code == 200
        assert client.post("/sessions", data={}).status_code == 403
    app.state.service.executor.shutdown(wait=True)


def test_hosted_agent_requires_server_token_and_disables_browser_key_setup(monkeypatch):
    monkeypatch.setenv("AGENT_HOSTED", "1")
    monkeypatch.delenv("AGENT_API_TOKEN", raising=False)
    with pytest.raises(RuntimeError, match="AGENT_API_TOKEN"):
        create_app(ScriptedModel)
    monkeypatch.setenv("AGENT_API_TOKEN", "explicit-test-only-service-token")
    app = create_app(ScriptedModel)
    with TestClient(app, headers=HEADERS) as client:
        assert client.get("/health").status_code == 200
        assert client.get("/status").status_code == 401
        assert client.get("/status", headers={"Authorization": "Bearer wrong-test-token"}).status_code == 401
        client.headers["Authorization"] = "Bearer explicit-test-only-service-token"
        assert client.get("/status").json()["canConfigure"] is False
        assert client.post("/configuration", json={"apiKey": "explicit-test-only-model-key"}).status_code == 403
    app.state.service.executor.shutdown(wait=True)


def test_failed_model_step_resumes_checkpoint_without_duplicate_answer(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)

    class FailOnce(ScriptedModel):
        failed = False

        def interpret(self, question, answer, proposal):
            if not self.failed:
                self.failed = True
                raise RuntimeError("Intentional test-only provider failure")
            return super().interpret(question, answer, proposal)

    app = create_app(FailOnce)
    with TestClient(app, headers=HEADERS) as client:
        session = start(client)
        session = reply(client, session, "Root cause answer.", "failure-reply-unique")
        assert session["status"] == "failed"
        response = client.post(f"/sessions/{session['sessionId']}/retry", json={"expectedRevision": session["revision"]})
        assert response.status_code == 202
        session = settled(client, session["sessionId"])
        assert session["status"] == "awaiting-input"
        assert session["confirmed"]["root_cause"] == "Root cause answer."
        assert sum(item["content"] == "Root cause answer." for item in session["messages"]) == 1
    app.state.service.executor.shutdown(wait=True)


def test_real_initial_interrupt_needs_no_model_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = create_app()
    with TestClient(app, headers=HEADERS) as client:
        assert client.get("/status").json()["ready"] is False
        session = start(client)
        assert session["status"] == "awaiting-input"
        assert "root cause" in session["message"].lower()
        assert session["artifacts"] == []
    app.state.service.executor.shutdown(wait=True)


def test_reconnected_key_is_used_by_existing_conversation(monkeypatch):
    import app as adapter

    calls = []

    class KeyAwareStandIn:
        def __init__(self):
            calls.append(adapter.os.getenv("ANTHROPIC_API_KEY"))
        name = "explicit-test-provider"

    monkeypatch.setattr(adapter, "ClaudeLLM", KeyAwareStandIn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "first-test-only-key")
    provider = adapter.LazyClaude()
    assert provider.name == "explicit-test-provider"
    assert provider.name == "explicit-test-provider"
    monkeypatch.setenv("ANTHROPIC_API_KEY", "reconnected-test-only-key")
    assert provider.name == "explicit-test-provider"
    assert calls == ["first-test-only-key", "reconnected-test-only-key"]


def test_actual_anonymized_deck_drives_latest_graph_and_review(monkeypatch):
    import os
    from presentation import DeckAnalysis, ProcessRisk, PresentationClaude, DeckGradeSuggestion
    fixture = os.getenv("REPORT_PPTX")
    if not fixture: pytest.skip("Set REPORT_PPTX for the private uploaded reference.")
    class PresentationModel(ScriptedModel):
        def suggest_grade(self, material):
            return PresentationClaude.suggest_grade(self, material)
        def _parsed(self, task, schema, material=None):
            import prompts
            assert task.startswith(prompts.SUGGEST_GRADE)
            assert schema is DeckGradeSuggestion
            return DeckGradeSuggestion(grade="B", reason="Explicit test rationale citing slide 8.", evidenceSlides=[8,999], missingEvidenceQuestion="")
        def analyze_presentation(self, material):
            assert "Finding title:" in material[0]["text"]
            assert "--- Slide 13 ---" in material[0]["text"]
            return DeckAnalysis(overview="Explicit test-only overview.", firstQuestion="Explicit tailored test-only root cause question?", processRisk=ProcessRisk(processName="Test process", grade=None, rationale="Pending grade step", evidenceSlides=[]))
        def write_summary(self, material, confirmed, tone, request, questions, answers):
            return BoardSummary(exe_summary="## Executive Board Summary\nExplicit test-only conclusion.\n\n## Potential Gaps for Executive Board Consideration\nExplicit private test gap.", pos_points=["Explicit test positive."])
    app = create_app(PresentationModel)
    targets = [
        {"id":"slide-4-table-3-0-0","label":"Audit conclusion","maxChars":1800},
        {"id":"slide-4-table-27-0-0","label":"Positive aspects","maxChars":650},
        {"id":"slide-4-table-5-1-0","label":"Main finding 1","maxChars":360},
        {"id":"slide-4-table-5-1-1","label":"Recommendation 1","maxChars":280},
    ]
    with TestClient(app, headers=HEADERS) as client:
        response=client.post("/sessions", data={"template":json.dumps({"name":"Private deck", "slides":[{}]*13}),"request":"Write compact report summary", "plan":json.dumps({"summaryFields":targets})}, files={"presentation":("Private report.pptx",Path(fixture).read_bytes())})
        assert response.status_code==202,response.text
        session=settled(client,response.json()["sessionId"])
        assert session["analysis"]["processRisk"]["grade"] is None
        assert "tailored test-only" in session["message"]
        for i, text in enumerate(["Cause", "Relationships", "Storyline", "Positives", "C", "Balanced", "The process owner"]):
            session=reply(client, session, text, f"actual-deck-reply-{i}")
        artifact=session["artifacts"][0]
        assert artifact["grade"]=="C"
        assert artifact["processRisk"]["grade"]=="B"
        assert artifact["processRisk"]["evidenceSlides"]==[8]
        assert "private test gap" not in artifact["fields"][targets[0]["id"]]
        assert "suggested by AI" in artifact["fields"][targets[3]["id"]]
        fields={**artifact["fields"],targets[0]["id"]:"Reviewed test conclusion."}
        response=client.post(f"/sessions/{session['sessionId']}/review",json={"expectedRevision":session["revision"],"content":artifact["content"],"fields":fields,"grade":"D"})
        assert response.status_code==200,response.text
        approved=response.json()["artifacts"][0]
        assert approved["reviewedGrade"]=="D"
        assert approved["processRisk"]["grade"]=="B"
        assert approved["processRisk"]["evidenceSlides"]==[8]
        assert approved["fields"][targets[0]["id"]]=="Reviewed test conclusion."
    app.state.service.executor.shutdown(wait=True)


def test_placeholder_deck_requests_grade_evidence_and_resumes_latest_graph():
    import io
    from pptx import Presentation
    from pptx.util import Inches
    from presentation import DeckAnalysis, ProcessRisk, PresentationClaude, DeckGradeSuggestion
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(8), Inches(.6))
    title.text = "Executive Summary"
    conclusion = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    conclusion.text = "Audit conclusion\n[AUDIT SUMMARY]\n[CURRENT MITIGATING CONTROLS]"
    source = io.BytesIO(); deck.save(source)
    class EvidenceModel(ScriptedModel):
        def analyze_presentation(self, material):
            return DeckAnalysis(overview="Explicit placeholder-deck test.", firstQuestion="What evidence is missing?", processRisk=ProcessRisk(processName="Test process", grade=None, rationale="Pending evidence", evidenceSlides=[]))
        def suggest_grade(self, material):
            return PresentationClaude.suggest_grade(self, material)
        def _parsed(self, task, schema, material=None):
            import prompts
            assert task.startswith(prompts.SUGGEST_GRADE)
            assert schema is DeckGradeSuggestion
            enough = any("A concrete test finding with insufficient controls" in item.get("text", "") for item in material)
            return DeckGradeSuggestion(grade="C" if enough else None, reason="Explicit test grade supported by slide 1 and audit-manager clarification." if enough else "The report contains placeholders without audit evidence.", evidenceSlides=[1,999] if enough else [], missingEvidenceQuestion="Describe the finding and current controls." if not enough else "")
        def write_summary(self, material, confirmed, tone, request, questions, answers):
            return BoardSummary(exe_summary="## Executive Board Summary\nExplicit test draft after clarification.", pos_points=[])
    app = create_app(EvidenceModel)
    target={"id":f"slide-1-shape-{conclusion.shape_id}","label":"Audit conclusion","maxChars":1800}
    with TestClient(app, headers=HEADERS) as client:
        response=client.post("/sessions",data={"template":json.dumps(TEMPLATE),"request":"Write test report", "plan":json.dumps({"summaryFields":[target]})},files={"presentation":("Placeholder-only test.pptx",source.getvalue())})
        assert response.status_code==202,response.text
        session=settled(client,response.json()["sessionId"])
        for i,text in enumerate(["Cause", "Relationships", "Storyline", "Positives"]):
            session=reply(client,session,text,f"placeholder-reply-{i}")
        assert session["status"]=="awaiting-input"
        assert session["analysis"]["processRisk"]["grade"] is None
        assert session["artifacts"]==[] and session["error"] is None
        assert "More evidence is needed" in session["message"]
        assert session["confirmed"]["positives"]=="Positives"
        session=reply(client,session,"I do not know yet.","evidence-unknown-unique")
        assert session["analysis"]["processRisk"]["grade"] is None
        assert "Describe the finding" in session["message"]
        session=reply(client,session,"A concrete test finding with insufficient controls","evidence-details-unique")
        assert session["status"]=="awaiting-input"
        assert session["analysis"]["processRisk"]["grade"]=="C"
        assert session["analysis"]["processRisk"]["evidenceSlides"]==[1]
        assert "grade" in session["message"].lower()
        for i,text in enumerate(["yes", "Balanced", "The process owner"]):
            session=reply(client,session,text,f"post-evidence-reply-{i}")
        assert session["status"]=="awaiting-review"
        assert session["confirmed"]["grade"]=="C"
        assert session["artifacts"][0]["processRisk"]["evidenceSlides"]==[1]
    app.state.service.executor.shutdown(wait=True)
