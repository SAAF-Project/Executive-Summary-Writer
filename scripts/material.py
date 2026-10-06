"""Loads the audit material (draft report, observations, management responses) as Claude content blocks."""

import base64
from pathlib import Path
from typing import Any, Dict, List

from sanitizer import sanitize_evidence

TEXT_SUFFIXES = {".txt", ".md"}



def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def load_material(paths: List[Path]) -> List[Dict[str, Any]]:
    """One content block per file: text for .txt / .md / .docx, a document block for .pdf."""
    blocks: List[Dict[str, Any]] = []
    for path in paths:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
            blocks.append(
                {
                    "type": "document",
                    "title": path.name,
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": data,
                    },
                }
            )
            continue
        if suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif suffix == ".docx":
            text = _docx_text(path)
        else:
            raise ValueError(
                f"Unsupported file type: {path.name} (use .txt, .md, .docx or .pdf)"
            )
        if not text.strip():
            raise ValueError(f"No text found in {path.name}")
        
        fenced_block, nonce, defanged = sanitize_evidence(text.strip(), tag_name="audit_material")
        # Include file metadata in the outer opening tag
        fenced_block = fenced_block.replace(
            f'<audit_material nonce="{nonce}">',
            f'<audit_material file="{path.name}" nonce="{nonce}">',
            1
        )
        blocks.append(
            {
                "type": "text",
                "text": fenced_block,
            }
        )
    return blocks

